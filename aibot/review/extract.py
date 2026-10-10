"""
Text extraction from uploaded PDF, DOCX and DOC files.
"""
import io
import os
import re
import subprocess  # nosec B404 - used only to run antiword with a fixed argument list
import tempfile

import pdfplumber
from docx import Document

ALLOWED_EXTENSIONS = {".pdf", ".doc", ".docs", ".docx"}

# Source files accepted by rubrics with `input: code`.
CODE_EXTENSIONS = frozenset({
    ".py", ".pyi", ".js", ".jsx", ".mjs", ".cjs", ".ts", ".tsx", ".vue", ".svelte",
    ".java", ".kt", ".kts", ".scala", ".groovy", ".gradle", ".go", ".rs", ".c", ".h",
    ".cpp", ".cc", ".cxx", ".hpp", ".cs", ".fs", ".vb", ".php", ".rb", ".swift", ".m",
    ".dart", ".lua", ".pl", ".r", ".jl", ".ex", ".exs", ".erl", ".hs", ".clj", ".ml",
    ".zig", ".sol", ".sql", ".sh", ".bash", ".zsh", ".ps1", ".bat", ".html", ".css",
    ".scss", ".less", ".json", ".yaml", ".yml", ".toml", ".xml", ".tf",
})
CODE_FILENAMES = frozenset({"dockerfile", "makefile", "jenkinsfile"})
# Starts each file when several are reviewed together; line numbers restart after it.
FILE_HEADER = re.compile(r"^==> (.+) <==$")


class UnsupportedDocument(ValueError):
    """The file type is not supported or its text could not be read."""


def normalize_whitespace(text: str) -> str:
    """Normalize whitespace for stable prompts and cache keys."""
    return re.sub(r"\s+", " ", text).strip()


def normalize_code(text: str) -> str:
    """Keep line breaks and indentation; drop trailing spaces and blank lines at the edges."""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    return "\n".join(line.rstrip() for line in lines).strip("\n")


def clean_filename(filename: str) -> str:
    """A file name (or relative path) safe to put in a file header."""
    return re.sub(r"[\x00-\x1f<>]", "", filename).strip()[:200]


def file_header(filename: str) -> str:
    return f"==> {filename} <=="


def is_code_file(filename: str) -> bool:
    """True for source files a code rubric accepts."""
    name = os.path.basename(filename.lower())
    return name in CODE_FILENAMES or os.path.splitext(name)[1] in CODE_EXTENSIONS


def extract_code_text(file_bytes: bytes) -> str:
    """Read a source file as UTF-8 text."""
    if b"\x00" in file_bytes:
        raise UnsupportedDocument("The file is not a text file.")
    try:
        return file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise UnsupportedDocument("Could not read the file as UTF-8 text.") from exc


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from each PDF page."""
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract paragraph text from DOCX files."""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    return "\n".join(paragraphs)


def extract_text_from_doc(file_bytes: bytes) -> str:
    """Use antiword for legacy DOC files."""
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp_file:  # nosec B108 - delete=False required so antiword can read the file by path
        tmp_file.write(file_bytes)
        tmp_path = tmp_file.name
    try:
        result = subprocess.run(  # nosec B603 B607 - fixed command list, no shell, no user-controlled input in args
            ["antiword", tmp_path],
            capture_output=True,
            text=True,
            check=False
        )
        if result.returncode != 0:
            raise UnsupportedDocument("Failed to extract text from DOC file.")
        return result.stdout
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def extract_document_text(file_bytes: bytes, filename: str) -> str:
    """Route extraction by file extension."""
    _, ext = os.path.splitext(filename.lower())
    extractors = {
        ".pdf": extract_text_from_pdf,
        ".docx": extract_text_from_docx,
        ".doc": extract_text_from_doc,
        ".docs": extract_text_from_doc,
    }
    if ext not in extractors:
        raise UnsupportedDocument("Unsupported file type.")
    try:
        return extractors[ext](file_bytes)
    except UnsupportedDocument:
        raise
    except Exception as exc:
        raise UnsupportedDocument("Could not read the file. Is it a valid document?") from exc
